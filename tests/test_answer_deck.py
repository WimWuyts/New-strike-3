"""Het correctiedeck moet klopppen zonder dat iemand het openklikt.

LibreOffice draait niet in elke omgeving, dus visuele QA is niet altijd
mogelijk. Deze tests controleren daarom op XML-niveau wat een leerkracht
anders pas in de klas zou merken: verdwenen antwoorden, tekst die van de slide
loopt, en animaties die naar niets wijzen.
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest

pytest.importorskip("pptx", reason="python-pptx niet geïnstalleerd")

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from lib.project import ROOT  # noqa: E402

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
BUILDER = ROOT / "scripts" / "build_answer_deck.py"
FIXTURE_THEME = "ace3-u2"


@pytest.fixture(scope="module")
def deck(tmp_path_factory) -> Presentation:
    key_path = ROOT / "data" / "answers" / f"{FIXTURE_THEME}.json"
    if not key_path.exists():
        subprocess.run(
            [sys.executable, str(ROOT / "scripts" / "make_test_fixture.py")],
            cwd=ROOT,
            check=True,
            capture_output=True,
        )

    out = tmp_path_factory.mktemp("deck") / "answers.pptx"
    result = subprocess.run(
        [sys.executable, str(BUILDER), "--theme", FIXTURE_THEME, "--out", str(out)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    return Presentation(out)


@pytest.fixture(scope="module")
def key() -> dict:
    return json.loads(
        (ROOT / "data" / "answers" / f"{FIXTURE_THEME}.json").read_text(encoding="utf-8")
    )


def _is_textbox(shape) -> bool:
    nv = shape._element.find(qn("p:nvSpPr"))
    props = nv.find(qn("p:cNvSpPr")) if nv is not None else None
    return props is not None and props.get("txBox") == "1"


def _all_text(deck: Presentation) -> str:
    return "\n".join(
        shape.text_frame.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


# ---------------------------------------------------------------------------
# Inhoud
# ---------------------------------------------------------------------------


def test_every_exercise_appears(deck: Presentation, key: dict):
    text = _all_text(deck)
    missing = [e["number"] for e in key["exercises"] if e["number"] not in text]
    assert missing == []


def test_no_answer_is_silently_dropped(deck: Presentation, key: dict):
    """De tabelbouwer liet ooit rijen vallen die niet meer pasten."""
    text = _all_text(deck)
    missing = []
    for exercise in key["exercises"]:
        if exercise.get("open_ended"):
            continue  # daar wordt de leerkrachtnotitie getoond, geen sleutel
        for item in exercise.get("items", []):
            if item["answer"] not in text:
                missing.append((exercise["number"], item["number"]))
        for row in (exercise.get("table") or {}).get("rows", []):
            for cell in row[1:]:
                if cell and cell not in text:
                    missing.append((exercise["number"], cell[:20]))
    assert missing == []


def test_open_exercises_show_the_teacher_note(deck: Presentation, key: dict):
    text = _all_text(deck)
    for exercise in key["exercises"]:
        if exercise.get("open_ended"):
            assert exercise["note_nl"][:40] in text
    assert "OPEN OPDRACHT" in text


def test_uncertain_answers_carry_a_warning(deck: Presentation, key: dict):
    has_uncertain = any(
        item.get("uncertain")
        for exercise in key["exercises"]
        for item in exercise.get("items", [])
    )
    if not has_uncertain:
        pytest.skip("De fixture bevat geen onzekere antwoorden.")
    assert "Niet met zekerheid leesbaar" in _all_text(deck)


# ---------------------------------------------------------------------------
# Vormgeving
# ---------------------------------------------------------------------------


def test_nothing_runs_off_the_slide(deck: Presentation):
    height, width = deck.slide_height, deck.slide_width
    for index, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if shape.top is None:
                continue
            assert shape.top + shape.height <= height, f"slide {index} loopt onderaan uit"
            assert shape.left + shape.width <= width + Inches(0.02), (
                f"slide {index} loopt rechts uit"
            )


def test_no_two_text_boxes_overlap(deck: Presentation):
    for index, slide in enumerate(deck.slides, 1):
        boxes = [
            (s.left, s.top, s.left + s.width, s.top + s.height, s.text_frame.text[:24])
            for s in slide.shapes
            if _is_textbox(s)
        ]
        for first in range(len(boxes)):
            for second in range(first + 1, len(boxes)):
                a, b = boxes[first], boxes[second]
                horizontal = min(a[2], b[2]) - max(a[0], b[0])
                vertical = min(a[3], b[3]) - max(a[1], b[1])
                assert not (horizontal > Inches(0.4) and vertical > Inches(0.10)), (
                    f"slide {index}: {a[4]!r} overlapt {b[4]!r}"
                )


# De body ligt tussen de kopregel en de voettekst. Daarbuiten staan labels
# (paginanummer, klik-hint, voetnoot) die bewust kleiner mogen zijn.
BODY_BAND = (Inches(1.5), Inches(6.65))


def test_body_text_is_at_least_22pt(deck: Presentation):
    """Achteraan de klas moet de inhoud leesbaar blijven."""
    for index, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or shape.top is None:
                continue
            if not (BODY_BAND[0] <= shape.top <= BODY_BAND[1]):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    text = (run.text or "").strip()
                    if not text or run.font.size is None:
                        continue
                    assert run.font.size >= Pt(22), (
                        f"slide {index}: {text[:20]!r} staat op {run.font.size.pt} pt"
                    )


def test_no_slide_is_empty(deck: Presentation):
    for index, slide in enumerate(deck.slides, 1):
        assert len(slide.shapes) > 0, f"slide {index} is leeg"


# ---------------------------------------------------------------------------
# Klik-animatie
# ---------------------------------------------------------------------------


def test_answer_slides_carry_a_timing_tree(deck: Presentation):
    animated = sum(
        1 for slide in deck.slides if slide._element.find(f"{{{P_NS}}}timing") is not None
    )
    assert animated > 0, "Geen enkele slide heeft klik-animatie."


def test_every_animation_targets_a_shape_on_its_own_slide(deck: Presentation):
    """Een animatie die naar een verdwenen shape wijst, breekt de presentatie."""
    for index, slide in enumerate(deck.slides, 1):
        timing = slide._element.find(f"{{{P_NS}}}timing")
        if timing is None:
            continue
        ids = {shape.shape_id for shape in slide.shapes}
        for target in timing.iter(f"{{{P_NS}}}spTgt"):
            assert int(target.get("spid")) in ids, (
                f"slide {index}: animatie wijst naar shape {target.get('spid')}"
            )


def test_each_step_waits_for_a_click(deck: Presentation):
    """Elke stap moet op een klik wachten, anders loopt alles ineens af."""
    for index, slide in enumerate(deck.slides, 1):
        timing = slide._element.find(f"{{{P_NS}}}timing")
        if timing is None:
            continue
        sequence = timing.find(f".//{{{P_NS}}}seq")
        steps = sequence.find(f"{{{P_NS}}}cTn").find(f"{{{P_NS}}}childTnLst")
        for step in steps:
            condition = step.find(f"{{{P_NS}}}cTn").find(f"{{{P_NS}}}stCondLst")
            delays = [c.get("delay") for c in condition]
            assert "indefinite" in delays, f"slide {index}: een stap wacht niet op een klik"


def test_open_exercise_slides_have_no_animation(deck: Presentation):
    """Bij een open opdracht valt niets te onthullen."""
    for slide in deck.slides:
        text = " ".join(
            s.text_frame.text for s in slide.shapes if s.has_text_frame
        )
        if "OPEN OPDRACHT" in text:
            assert slide._element.find(f"{{{P_NS}}}timing") is None
